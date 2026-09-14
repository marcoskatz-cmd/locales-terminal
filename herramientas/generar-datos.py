# -*- coding: utf-8 -*-
"""Convierte la planilla de relevamiento en los datos que consume el sistema.

  Entrada : docs/Relevamiento locales Terminal.xlsx   (hojas LOCALES, CONTRATOS, MANTENIMIENTO, LISTAS)
  Salidas : mock/datos.json        -> lo lee la app cuando CONFIG.API_URL está vacío (modo demo)
            backend/MockData.js    -> lo usa setupApp() para sembrar la planilla de Google

Uso: python herramientas/generar-datos.py

Reglas (mismas que la app):
  - Local con fila en CONTRATOS -> contrato VIGENTE y estado ALQUILADO. Si faltan datos del contrato
    (fechas, monto) igual se carga, con observación "INCOMPLETO", para que se vea en la app qué falta.
  - se_alquila = NO -> activo = false (no cuenta para ocupación ni deuda).
  - estado vacío -> LIBRE si no tiene contrato.
  - Si falta una columna o una hoja, se avisa y NO se genera nada (no fallar en silencio).
"""
import json, os, sys, datetime, re
from openpyxl import load_workbook

try: sys.stdout.reconfigure(encoding='utf-8')
except Exception: pass

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(RAIZ, 'docs', 'Relevamiento locales Terminal.xlsx')
ESPERADAS = {
    'LOCALES': ['id_local', 'etiqueta_plano', 'sector', 'nombre_en_plano', 'estado_deducido', 'nombre', 'm2', 'rubro', 'estado', 'se_alquila', 'observaciones'],
    'CONTRATOS': ['id_local', 'inquilino', 'cuit', 'contacto', 'fecha_inicio', 'fecha_fin', 'monto_alquiler', 'indice_ajuste', 'periodicidad_meses',
                  'proxima_fecha_ajuste', 'deposito', 'expensas_mensuales', 'observaciones'],
    'MANTENIMIENTO': ['id_local', 'fecha_reporte', 'tipo_falla', 'descripcion', 'prioridad', 'estado', 'quien_intervino', 'fecha_resolucion', 'costo', 'observaciones'],
    'LISTAS': ['rubros', 'sectores', 'indices_ajuste', 'tipos_falla', 'intervinientes', 'medios_pago'],
}

def fecha(v):
    if v in (None, ''): return ''
    if isinstance(v, (datetime.datetime, datetime.date)): return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', s)
    if m: return f'{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
    if re.match(r'^\d{4}-\d{2}-\d{2}$', s): return s
    raise ValueError(f'Fecha no reconocida: "{s}" (usar dd/mm/aaaa)')
def num(v, defecto=0):
    if v in (None, ''): return defecto
    if isinstance(v, (int, float)): return v
    s = str(v).replace('$', '').replace('.', '').replace(',', '.').strip()
    try: return float(s) if '.' in s else int(s)
    except ValueError: raise ValueError(f'Número no reconocido: "{v}"')
def txt(v): return '' if v is None else str(v).strip()

wb = load_workbook(XLSX, data_only=True)
problemas = []
hojas = {}
for nombre, cols in ESPERADAS.items():
    if nombre not in wb.sheetnames: problemas.append(f'Falta la hoja {nombre}'); continue
    ws = wb[nombre]
    enc = [txt(c.value) for c in ws[1]]
    faltan = [c for c in cols if c not in enc]
    if faltan: problemas.append(f'Hoja {nombre}: faltan columnas {faltan}')
    filas = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if all(v in (None, '') for v in r): continue
        filas.append({enc[j]: r[j] for j in range(min(len(enc), len(r))) if enc[j]})
    hojas[nombre] = filas
if problemas:
    print('NO SE GENERÓ NADA. Problemas en la planilla:'); [print(' -', p) for p in problemas]; sys.exit(1)

listas = {}
for k in ESPERADAS['LISTAS']:
    listas[k] = [txt(f.get(k)) for f in hojas['LISTAS'] if txt(f.get(k))]

# ---------- LOCALES ----------
locales, ids = [], set()
for f in hojas['LOCALES']:
    idl = txt(f['id_local'])
    if not idl: continue
    if idl in ids: problemas.append(f'LOCALES: id_local repetido {idl}'); continue
    ids.add(idl)
    locales.append({
        'id_local': idl, 'nombre': txt(f.get('nombre')) or 'Local ' + txt(f.get('etiqueta_plano')), 'planta': 'PB',
        'sector': txt(f.get('sector')), 'm2': num(f.get('m2'), 0), 'rubro': txt(f.get('rubro')),
        'estado': txt(f.get('estado')).upper(), 'observaciones': txt(f.get('observaciones')),
        'activo': txt(f.get('se_alquila')).upper() != 'NO',
    })

# ---------- CONTRATOS ----------
contratos, con_contrato = [], set()
for n, f in enumerate(hojas['CONTRATOS'], 1):
    idl = txt(f['id_local'])
    if not idl: continue
    if idl not in ids: problemas.append(f'CONTRATOS fila {n + 1}: el local {idl} no existe en LOCALES'); continue
    if idl in con_contrato: problemas.append(f'CONTRATOS: el local {idl} tiene más de un contrato'); continue
    con_contrato.add(idl)
    try:
        c = {'id_contrato': f'C-{len(contratos) + 1:03d}', 'id_local': idl, 'inquilino': txt(f.get('inquilino')), 'cuit': txt(f.get('cuit')),
             'contacto': txt(f.get('contacto')), 'fecha_inicio': fecha(f.get('fecha_inicio')), 'fecha_fin': fecha(f.get('fecha_fin')),
             'monto_alquiler': num(f.get('monto_alquiler')), 'indice_ajuste': txt(f.get('indice_ajuste')) or 'IPC',
             'periodicidad_meses': int(num(f.get('periodicidad_meses'), 3)), 'proxima_fecha_ajuste': fecha(f.get('proxima_fecha_ajuste')),
             'deposito': num(f.get('deposito')), 'expensas_mensuales': num(f.get('expensas_mensuales')),
             'estado_contrato': 'VIGENTE', 'observaciones': txt(f.get('observaciones'))}
    except ValueError as e:
        problemas.append(f'CONTRATOS fila {n + 1} ({idl}): {e}'); continue
    if not c['inquilino']: problemas.append(f'CONTRATOS fila {n + 1} ({idl}): falta el inquilino'); continue
    faltantes = [k for k in ('fecha_inicio', 'fecha_fin', 'monto_alquiler') if not c[k]]
    if faltantes:
        c['observaciones'] = ('INCOMPLETO: falta ' + ', '.join(faltantes) + '. ' + c['observaciones']).strip()
    contratos.append(c)

# Coherencia local <-> contrato (misma regla que la app)
for l in locales:
    if l['id_local'] in con_contrato:
        if l['estado'] != 'ALQUILADO': l['estado'] = 'ALQUILADO'
    else:
        if l['estado'] == 'ALQUILADO': problemas.append(f'LOCALES: {l["id_local"]} figura ALQUILADO pero no tiene fila en CONTRATOS'); l['estado'] = 'LIBRE'
        if not l['estado']: l['estado'] = 'LIBRE'
    if l['estado'] not in ('LIBRE', 'ALQUILADO', 'REFACCION'): problemas.append(f'LOCALES: estado inválido "{l["estado"]}" en {l["id_local"]}'); l['estado'] = 'LIBRE'

# ---------- MANTENIMIENTO ----------
mant = []
for n, f in enumerate(hojas['MANTENIMIENTO'], 1):
    idl = txt(f['id_local'])
    if not idl: continue
    if idl != 'COMUN' and idl not in ids: problemas.append(f'MANTENIMIENTO fila {n + 1}: el local {idl} no existe'); continue
    try:
        m = {'id_mant': f'M-{len(mant) + 1:03d}', 'id_local': idl, 'planta': 'PB', 'fecha_reporte': fecha(f.get('fecha_reporte')) or datetime.date.today().isoformat(),
             'tipo_falla': txt(f.get('tipo_falla')) or 'OTRO', 'descripcion': txt(f.get('descripcion')), 'prioridad': txt(f.get('prioridad')).upper() or 'MEDIA',
             'estado': txt(f.get('estado')).upper() or 'PENDIENTE', 'quien_intervino': txt(f.get('quien_intervino')), 'fecha_resolucion': fecha(f.get('fecha_resolucion')),
             'costo': num(f.get('costo'), ''), 'fotos': '', 'observaciones': txt(f.get('observaciones'))}
    except ValueError as e:
        problemas.append(f'MANTENIMIENTO fila {n + 1}: {e}'); continue
    mant.append(m)

if problemas:
    print('AVISOS (se generó igual, pero conviene corregir):'); [print(' -', p) for p in problemas]

datos = {'generado': datetime.date.today().isoformat(), 'origen': os.path.basename(XLSX), 'listas': listas, 'locales': locales, 'contratos': contratos,
         'cuotas': [], 'pagos': [], 'mantenimiento': mant, 'historial': []}
with open(os.path.join(RAIZ, 'mock', 'datos.json'), 'w', encoding='utf-8') as fh: json.dump(datos, fh, ensure_ascii=False, indent=1)
with open(os.path.join(RAIZ, 'backend', 'MockData.js'), 'w', encoding='utf-8') as fh:
    fh.write('// GENERADO por herramientas/generar-datos.py desde "' + os.path.basename(XLSX) + '" — NO editar a mano.\n'
             '// Lo usa setupApp() para sembrar la planilla de Google. Regenerar cada vez que cambie el relevamiento.\n'
             'const MOCK = ' + json.dumps(datos, ensure_ascii=False, indent=1) + ';\n')
incompletos = sum(1 for c in contratos if c['observaciones'].startswith('INCOMPLETO'))
print(f'OK → mock/datos.json y backend/MockData.js\n  locales={len(locales)} (activos {sum(1 for l in locales if l["activo"])}) · contratos={len(contratos)} '
      f'(incompletos {incompletos}) · mantenimiento={len(mant)}')

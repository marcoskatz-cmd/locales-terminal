# -*- coding: utf-8 -*-
"""Genera la planilla de relevamiento (Excel) prellenada con lo que se pudo leer del plano.
Es el archivo que administración completa con los datos reales de locales y contratos.

  Entrada : herramientas/plano-locales.json   (lo produce extraer-plano.py)
  Salida  : docs/Relevamiento locales Terminal.xlsx

Uso: python herramientas/generar-relevamiento.py
Si el archivo ya existe NO se pisa (para no perder lo cargado): borrarlo a mano o pasar --forzar.

Después de completarla, generar-datos.py la convierte en mock/datos.json y backend/MockData.js.
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

try: sys.stdout.reconfigure(encoding='utf-8')
except Exception: pass

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, 'docs', 'Relevamiento locales Terminal.xlsx')
if os.path.exists(SALIDA) and '--forzar' not in sys.argv:
    print(f'Ya existe {SALIDA}. No se pisa. Usá --forzar si querés regenerarla desde cero.'); sys.exit(0)

plano = json.load(open(os.path.join(RAIZ, 'herramientas', 'plano-locales.json'), encoding='utf-8'))
locales = plano['locales']

LISTAS = {
    'rubros': ['Boletería', 'Encomiendas / cargas', 'Kiosco', 'Cafetería', 'Panadería', 'Heladería', 'Comidas rápidas', 'Restaurante',
               'Farmacia', 'Indumentaria', 'Calzado', 'Lencería', 'Librería', 'Regalería', 'Juguetería', 'Perfumería', 'Telefonía',
               'Banco / cajero', 'Casino', 'Barbería / peluquería', 'Organismo público', 'Depósito', 'Servicios', 'Otro'],
    'sectores': ['Block 1', 'Block 2', 'Block 3', 'Block 4', 'Block 5', 'Block 6', 'Block 7', 'Block 8'],
    'indices_ajuste': ['IPC', 'ICL', '% FIJO', 'OTRO'],
    'tipos_falla': ['ELÉCTRICA', 'SANITARIA', 'ESTRUCTURAL', 'PINTURA', 'CERRAJERÍA', 'CLIMATIZACIÓN', 'VIDRIOS / CARPINTERÍA', 'OTRO'],
    'intervinientes': ['Personal INGECO', 'Electricista', 'Plomero', 'Albañil', 'Refrigeración', 'Cerrajero', 'Otro'],
    'medios_pago': ['TRANSFERENCIA', 'EFECTIVO', 'CHEQUE', 'DÉBITO AUTOMÁTICO'],
}
SOSPECHA_NO_ALQUILA = ('DEPÓSITO', 'DEPOSITO', 'POLICIA', 'POLICÍA', 'GENDARM', 'CNRT', 'COMISIÓN', 'COMISION', 'DESTAC', 'ENFERMER', 'MANTENIMIENTO')
def rubro_sugerido(nombre, idl):
    n = nombre.upper()
    if any(k in n for k in ('DEPÓSITO', 'DEPOSITO')): return 'Depósito'
    if any(k in n for k in ('POLICIA', 'POLICÍA', 'GENDARM', 'CNRT', 'COMISIÓN', 'COMISION', 'DESTAC', 'ENFERMER')): return 'Organismo público'
    if 'CARGAS' in n or 'ENCOM' in n or 'TRANSFER' in n: return 'Encomiendas / cargas'
    if 'FARMACIA' in n: return 'Farmacia'
    if 'PANADER' in n: return 'Panadería'
    if 'HELADO' in n: return 'Heladería'
    if 'KIOSCO' in n: return 'Kiosco'
    if 'BANCO' in n: return 'Banco / cajero'
    if 'CASINO' in n: return 'Casino'
    if 'RESTAUR' in n: return 'Restaurante'
    if 'BARBE' in n: return 'Barbería / peluquería'
    if 'LENCER' in n: return 'Lencería'
    if 'TELEF' in n or 'TLEFON' in n: return 'Telefonía'
    return ''

# ---------- estilos ----------
FUENTE_TIT = Font(bold=True, color='FFFFFF'); RELLENO_TIT = PatternFill('solid', fgColor='0F766E')
RELLENO_PLANO = PatternFill('solid', fgColor='E2E8F0')     # columnas que vienen del plano (no tocar)
RELLENO_CARGAR = PatternFill('solid', fgColor='FFFFFF')
RELLENO_OJO = PatternFill('solid', fgColor='FEF3C7')       # revisar
BORDE = Border(*(Side(style='thin', color='CBD5E1'),) * 4)

def encabezar(ws, cols, anchos, fila=1):
    for j, c in enumerate(cols, 1):
        cel = ws.cell(row=fila, column=j, value=c); cel.font = FUENTE_TIT; cel.fill = RELLENO_TIT
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); cel.border = BORDE
        ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[fila].height = 32
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)

def validar(ws, col_letra, lista_nombre, desde, hasta, lista_ws='LISTAS'):
    col = list(LISTAS.keys()).index(lista_nombre) + 1
    n = len(LISTAS[lista_nombre]) + 30
    dv = DataValidation(type='list', formula1=f"='{lista_ws}'!${get_column_letter(col)}$2:${get_column_letter(col)}${n + 1}", allow_blank=True, showErrorMessage=True,
                        errorTitle='Valor no permitido', error='Elegí una opción de la lista (se pueden agregar opciones en la hoja LISTAS).')
    ws.add_data_validation(dv); dv.add(f'{col_letra}{desde}:{col_letra}{hasta}')

def validar_fija(ws, col_letra, opciones, desde, hasta):
    dv = DataValidation(type='list', formula1='"' + ','.join(opciones) + '"', allow_blank=True, showErrorMessage=True,
                        errorTitle='Valor no permitido', error='Elegí una opción de la lista.')
    ws.add_data_validation(dv); dv.add(f'{col_letra}{desde}:{col_letra}{hasta}')

wb = Workbook()

# ---------- INSTRUCCIONES ----------
ws = wb.active; ws.title = 'INSTRUCCIONES'
ws.column_dimensions['A'].width = 120
lineas = [
    ('Relevamiento de locales · Terminal de Ómnibus · INGECO S.A.', True),
    ('', False),
    ('Esta planilla alimenta el sistema de gestión de locales. Se completa una vez; después los cambios se hacen desde la app.', False),
    ('', False),
    ('CÓMO ESTÁ ARMADA', True),
    ('• Hoja LOCALES: una fila por local detectado en el plano "CONFORME A OBRA 2026". Las columnas grises vienen del plano y no hay que tocarlas.', False),
    ('  Completar: nombre (si se lo quiere llamar distinto de "Local 23-25"), m², rubro, estado y si se alquila.', False),
    ('• Hoja CONTRATOS: una fila por local que en el plano figura con un inquilino. El nombre del inquilino ya está cargado; hay que completar el resto.', False),
    ('  Si hay un local alquilado que no figura, agregar una fila con su id_local. Si alguno de los listados ya no está alquilado, borrar la fila y poner el local en LIBRE.', False),
    ('• Hoja MANTENIMIENTO: opcional. Sirve para cargar fallas pendientes o el historial reciente. Se puede dejar vacía y cargar todo desde la app.', False),
    ('• Hoja LISTAS: las opciones de los desplegables. Se pueden agregar valores al final de cada columna (no borrar las columnas).', False),
    ('', False),
    ('REGLAS', True),
    ('• Los desplegables no aceptan texto libre: si falta una opción, agregarla en LISTAS.', False),
    ('• Fechas: escribirlas como fecha de Excel (dd/mm/aaaa). Montos: número sin puntos ni "$".', False),
    ('• estado: LIBRE / ALQUILADO / REFACCION. Un local con fila en CONTRATOS tiene que estar ALQUILADO, y viceversa.', False),
    ('• se_alquila = NO para depósitos propios, oficinas de organismos y espacios de uso interno: quedan en el mapa pero no cuentan para ocupación ni deuda.', False),
    ('• Filas en amarillo: hay algo para verificar (etiqueta repetida en el plano, posible organismo/depósito, o sin rótulo en el plano).', False),
    ('', False),
    ('QUÉ PASA DESPUÉS', True),
    ('Con la planilla completa se corre herramientas/generar-datos.py y la app pasa a mostrar los datos reales. Las cuotas de alquiler/expensas se generan desde la app, mes a mes.', False),
]
for i, (t, negrita) in enumerate(lineas, 1):
    c = ws.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=True, vertical='top')
    if negrita: c.font = Font(bold=True, size=12 if i == 1 else 11)

# ---------- LOCALES ----------
ws = wb.create_sheet('LOCALES')
cols = ['id_local', 'etiqueta_plano', 'sector', 'nombre_en_plano', 'estado_deducido', 'nombre', 'm2', 'rubro', 'estado', 'se_alquila', 'observaciones']
encabezar(ws, cols, [12, 14, 10, 26, 14, 18, 8, 22, 13, 11, 40])
ws.cell(row=1, column=6).comment = Comment('Cómo se ve el local en la app. Por defecto "Local <etiqueta>".', 'sistema')
ws.cell(row=1, column=10).comment = Comment('SI = local comercial que se alquila. NO = uso interno / organismo (no cuenta para ocupación).', 'sistema')
orden_sector = {s: i for i, s in enumerate(LISTAS['sectores'])}
locales_ord = sorted(locales, key=lambda l: (orden_sector.get(l['sector_nombre'], 99), l['centro'][1] > 415, l['centro'][0]))
for i, l in enumerate(locales_ord, 2):
    nombre_plano = l['nombre_en_plano']
    est = {'ALQUILADO': 'ALQUILADO', 'LIBRE': 'LIBRE'}.get(l['estado_deducido'], '')
    sospecha = any(k in nombre_plano.upper() for k in SOSPECHA_NO_ALQUILA)
    obs = []
    if l['estado_deducido'] == 'A CONFIRMAR': obs.append('Sin rótulo en el plano: confirmar estado')
    if '_' in l['id_local']: obs.append('Etiqueta repetida en el plano (hay dos "' + l['etiqueta_plano'] + '")')
    if sospecha: obs.append('¿Se alquila? Parece depósito/organismo')
    fila = [l['id_local'], l['etiqueta_plano'], l['sector_nombre'], nombre_plano, l['estado_deducido'],
            'Local ' + l['etiqueta_plano'], None, rubro_sugerido(nombre_plano, l['id_local']), est, 'SI', ' · '.join(obs)]
    for j, v in enumerate(fila, 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BORDE
        c.fill = RELLENO_PLANO if j <= 5 else RELLENO_CARGAR
        if obs and j > 5: c.fill = RELLENO_OJO
n = len(locales_ord) + 1
validar(ws, 'C', 'sectores', 2, n + 200)
validar(ws, 'H', 'rubros', 2, n + 200)
validar_fija(ws, 'I', ['LIBRE', 'ALQUILADO', 'REFACCION'], 2, n + 200)
validar_fija(ws, 'J', ['SI', 'NO'], 2, n + 200)
ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{n}'

# ---------- CONTRATOS ----------
ws = wb.create_sheet('CONTRATOS')
cols = ['id_local', 'inquilino', 'cuit', 'contacto', 'fecha_inicio', 'fecha_fin', 'monto_alquiler', 'indice_ajuste', 'periodicidad_meses',
        'proxima_fecha_ajuste', 'deposito', 'expensas_mensuales', 'observaciones']
encabezar(ws, cols, [12, 30, 15, 28, 13, 13, 15, 13, 12, 14, 14, 14, 40])
ws.cell(row=1, column=9).comment = Comment('Cada cuántos meses se ajusta el alquiler (1, 2, 3, 4, 6, 12).', 'sistema')
ws.cell(row=1, column=11).comment = Comment('Depósito en garantía entregado al firmar.', 'sistema')
i = 2
for l in locales_ord:
    if l['estado_deducido'] != 'ALQUILADO': continue
    fila = [l['id_local'], l['nombre_en_plano'].title() if l['nombre_en_plano'].isupper() else l['nombre_en_plano'], None, None, None, None, None, 'IPC', 3, None, None, None, None]
    for j, v in enumerate(fila, 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BORDE; c.fill = RELLENO_PLANO if j <= 2 else RELLENO_CARGAR
    for j in (5, 6, 10): ws.cell(row=i, column=j).number_format = 'DD/MM/YYYY'
    for j in (7, 11, 12): ws.cell(row=i, column=j).number_format = '#,##0'
    i += 1
validar(ws, 'H', 'indices_ajuste', 2, i + 200)
validar_fija(ws, 'I', ['1', '2', '3', '4', '6', '12'], 2, i + 200)
ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{i - 1}'

# ---------- MANTENIMIENTO ----------
ws = wb.create_sheet('MANTENIMIENTO')
cols = ['id_local', 'fecha_reporte', 'tipo_falla', 'descripcion', 'prioridad', 'estado', 'quien_intervino', 'fecha_resolucion', 'costo', 'observaciones']
encabezar(ws, cols, [12, 13, 20, 40, 11, 12, 20, 14, 12, 40])
ws.cell(row=1, column=1).comment = Comment('id_local de la hoja LOCALES, o COMUN para áreas comunes.', 'sistema')
validar(ws, 'C', 'tipos_falla', 2, 500)
validar_fija(ws, 'E', ['ALTA', 'MEDIA', 'BAJA'], 2, 500)
validar_fija(ws, 'F', ['PENDIENTE', 'EN CURSO', 'RESUELTO'], 2, 500)
validar(ws, 'G', 'intervinientes', 2, 500)
for r in range(2, 500):
    for j in (2, 8): ws.cell(row=r, column=j).number_format = 'DD/MM/YYYY'

# ---------- LISTAS ----------
ws = wb.create_sheet('LISTAS')
encabezar(ws, list(LISTAS.keys()), [24, 12, 16, 24, 20, 20])
for j, k in enumerate(LISTAS.keys(), 1):
    for i, v in enumerate(LISTAS[k], 2): ws.cell(row=i, column=j, value=v).border = BORDE

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
wb.save(SALIDA)
alq = sum(1 for l in locales if l['estado_deducido'] == 'ALQUILADO')
print(f'OK → {SALIDA}\n  LOCALES: {len(locales)} filas · CONTRATOS: {alq} filas prellenadas con el inquilino del plano')

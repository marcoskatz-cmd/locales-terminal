# -*- coding: utf-8 -*-
"""Arma la planilla de relevamiento cruzando las TRES fuentes reales:

  1. "Alquileres Terminal.xlsx" (hoja Locales Comerciales + Hoja 11)   -> lista maestra de UNIDADES
     (boleterías, locales, depósitos, góndolas, oficinas, encomiendas, predio norte), superficie, rubro,
     inquilino, precio acordado, expensas, notas ("se va", "Gestion Judicial", ...).
  2. "COBRANZAS mayo-Agosto ....xlsx" (hojas MAYO/JUNIO/JULIO/AGOSTO)   -> CUOTAS del mes (TOTAL) y PAGOS
     (transferencia / efectivo / cheque / retenciones, con fecha).
  3. herramientas/plano-locales.json (de extraer-plano.py)              -> zonas del plano (id L-XX).

  Salida: docs/Relevamiento locales Terminal.xlsx con hojas LOCALES, CONTRATOS, CUOTAS, PAGOS, LISTAS,
          CRUCE (informe del cruce: qué se emparejó con qué y con qué confianza) y REVISAR (dudas).
          Después, generar-datos.py convierte esa planilla en los datos de la app.

Uso:
  python herramientas/importar-alquileres.py "C:\\...\\Alquileres Terminal.xlsx" "C:\\...\\COBRANZAS ....xlsx"

Regla general: NO se descarta nada. Lo que no se puede emparejar con certeza igual entra, marcado en la hoja
REVISAR, para que administración lo confirme o corrija en la planilla. Los ids de las unidades se arman por
categoría + número (BOL-1-2, LOC-501, GON-3, OFI-7, DEP-305, ENC-512...) porque el número solo se repite entre
categorías (hay un "9" boletería, un "9" local y un "9" góndola).
"""
import json, os, re, sys, unicodedata, datetime
from difflib import SequenceMatcher
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

try: sys.stdout.reconfigure(encoding='utf-8')
except Exception: pass

if len(sys.argv) < 3:
    print(__doc__); sys.exit(1)
XLS_ALQ, XLS_COB = sys.argv[1], sys.argv[2]
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SALIDA = os.path.join(RAIZ, 'docs', 'Relevamiento locales Terminal.xlsx')
PLANO = json.load(open(os.path.join(RAIZ, 'herramientas', 'plano-locales.json'), encoding='utf-8'))
HOY = datetime.date.today()

# ============================================================ utilidades
def sin_acentos(s): return ''.join(c for c in unicodedata.normalize('NFD', str(s)) if unicodedata.category(c) != 'Mn')
def txt(v): return '' if v is None else str(v).strip()
def num(v):
    if v in (None, ''): return 0
    if isinstance(v, (int, float)): return float(v)
    s = str(v).replace('$', '').replace('.', '').replace(',', '.').strip()
    try: return float(s)
    except ValueError: return 0
def es_num(v): return isinstance(v, (int, float)) or (isinstance(v, str) and re.match(r'^-?\d+([.,]\d+)?$', v.strip() or 'x'))
def fecha(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    return ''
def limpiar_id(v):
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return txt(v)

SUFIJOS = r'\b(S\.?\s?R\.?\s?L\.?|S\.?\s?A\.?\s?S\.?|S\.?\s?A\.?|SACIFI|S A C I F E I|LTDA|SOCIEDAD ANONIMA|SRL|SAS|SA|HNOS|HERMANOS|DE|DEL|LA|EL|LOS|Y|E)\b'
def norm_nombre(s):
    s = sin_acentos(txt(s)).upper().replace('ﾑ', 'N').replace('`', "'")
    s = re.sub(r"[^A-Z0-9 ']", ' ', s)
    s = re.sub(SUFIJOS, ' ', s)
    s = re.sub(r"'S?\b", '', s)
    return ' '.join(s.split())
def similitud(a, b):
    a, b = norm_nombre(a), norm_nombre(b)
    if not a or not b: return 0
    wa, wb = set(a.split()), set(b.split())
    jac = len(wa & wb) / len(wa | wb) if wa | wb else 0
    return max(jac, SequenceMatcher(None, a, b).ratio())

def tokens(texto):
    """'48- 49- 50- 51- 52- 53' -> {48..53} · '41 A - 41 B' -> {41A,41B} · '53 a 44' -> {44..53} · '10a&b' -> {10A,10B}"""
    t = sin_acentos(txt(texto)).upper()
    t = re.sub(r'\bBLOCK\s+[IVX\d]+\s*PA\b', ' ', t)
    t = re.sub(r'\b(LOCAL|OFICINA|LOCALES)\b', ' ', t)
    t = re.sub(r'\((.*?)\)', r' \1 ', t)
    t = re.sub(r'(\d+)\s*([A-D])\s*(?:&|Y)\s*([A-D])\b', r'\1\2 \1\3', t)     # 10a&b, 518C y D
    t = re.sub(r'(\d+)\s+([A-D])\b(?!\s*\d)', r'\1\2', t)                      # 41 A -> 41A (pero '53 a 44' sigue siendo un rango)
    partes = list(re.finditer(r'[A-Z]{0,2}\d+[A-D]?', t))
    toks, prev = [], None
    for m in partes:
        tok = m.group(0)
        if prev is not None:
            sep = t[prev.end():m.start()]
            a, b = prev.group(0), tok
            if a.isdigit() and b.isdigit() and re.match(r'^\s*(-|/|A|AL)\s*$', sep):
                lo, hi = sorted((int(a), int(b)))
                if 1 <= hi - lo <= 10:
                    toks += [str(x) for x in range(lo, hi + 1)]
        if tok not in toks: toks.append(tok)
        prev = m
    return set(toks)

def compactar(toks):
    ints = sorted(int(x) for x in toks if x.isdigit()); otros = sorted(x for x in toks if not x.isdigit())
    partes = []
    if ints:
        if len(ints) > 2 and ints == list(range(ints[0], ints[-1] + 1)): partes.append(f'{ints[0]}-{ints[-1]}')
        else: partes.append('-'.join(str(i) for i in ints))
    partes += otros
    return '-'.join(partes)

def slug(s): return re.sub(r'[^A-Z0-9]+', '-', sin_acentos(s).upper()).strip('-')

# ============================================================ 1. Unidades (Alquileres Terminal)
CATEGORIAS = {'boleteria': ('BOL', 'Boletería'), 'local comercial': ('LOC', 'Local comercial'), 'deposito': ('DEP', 'Depósito'),
              'gondola': ('GON', 'Góndola'), 'oficinas': ('OFI', 'Oficina'), 'encomiendas': ('ENC', 'Encomiendas'),
              'corta distancia': ('CD', 'Corta distancia')}
INTERNOS = ('TERMINAL', 'ADMINISTRACION', 'DIRECCION DE TRANSPORTE', 'DEPOSITOS TERMINAL', 'TALLER', 'ENFERMERIA', 'LIMPIEZA',
            'OPERACIONES', 'SEGURIDAD', 'POLICIA', 'CNRT', 'GENDARMERIA', 'OFICINA DAU', 'TURISMO')
EXTERIORES = ('PN1', 'PN2', 'J1', 'S1', 'CARTEL', 'ESTACION', 'TAXIS')

wb_alq = load_workbook(XLS_ALQ, data_only=True)
ws = wb_alq['Locales Comerciales']
enc = [txt(c.value) for c in ws[1]]
COL = {'id': 0, 'flag': 1, 'block': 2, 'inquilino': 3, 'rubro': 4, 'm2': 5, 'entrepiso': 6, 'vidriera': 7, 'precio_viejo': 8, 'nuevo': 9,
       'final': 10, 'contra': 14, 'respuesta': 15, 'sigue': 16, 'categoria': 17, 'notas': 18}
# Chequeo de columnas: si cambió el orden, avisar y parar
esperado = {'block': 'Block', 'rubro': 'Rubro', 'm2': 'Superficie', 'final': 'Precio final', 'respuesta': 'Respuesta (PRECIO FINAL)', 'sigue': 'Sigue?', 'categoria': 'Categoria', 'notas': 'Notas'}
for k, nombre in esperado.items():
    if not enc[COL[k]].startswith(nombre): print(f'ATENCIÓN: en "Locales Comerciales" esperaba "{nombre}" en la columna {COL[k] + 1} y encontré "{enc[COL[k]]}". Reviso el mapeo antes de seguir.'); sys.exit(1)

# Expensas (Hoja 11): por (id normalizado, block)
expensas = {}
if 'Hoja 11' in wb_alq.sheetnames:
    h = wb_alq['Hoja 11']; e11 = [txt(c.value) for c in h[1]]
    if 'Expensas' in e11:
        ce = e11.index('Expensas')
        for r in h.iter_rows(min_row=2, values_only=True):
            if r[0] in (None, '') or not es_num(r[1]): continue
            expensas[(compactar(tokens(limpiar_id(r[0]))), int(num(r[1])))] = num(r[ce])

CORRECCIONES_ID = {'59 - 50': '59 - 60'}   # errores de tipeo evidentes en la planilla de alquileres
unidades, ids_usados = [], {}
for r in ws.iter_rows(min_row=2, values_only=True):
    idr, blk, inq, cat = limpiar_id(r[COL['id']]), r[COL['block']], txt(r[COL['inquilino']]), txt(r[COL['categoria']])
    idr = CORRECCIONES_ID.get(idr, idr)
    if not es_num(blk) or not (1 <= int(num(blk)) <= 8): continue
    if not re.search(r'[A-Za-z]', inq) or es_num(inq): continue
    if not idr: continue
    catk = sin_acentos(cat).lower().strip()
    pref, catn = CATEGORIAS.get(catk, ('OTR', cat or 'Otro'))
    toks = tokens(idr)
    inq_n = sin_acentos(inq).upper()
    interno = any(k in inq_n for k in INTERNOS) and 'PRIVAD' not in inq_n
    exterior = any(k in sin_acentos(idr).upper() for k in EXTERIORES) or 'PREDIO' in sin_acentos(idr).upper()
    if toks and not re.search(r'[A-Z]{3,}', sin_acentos(idr).upper().replace('OFICINA', '').replace('LOCAL', '').replace('PA', '')):
        base = f'{pref}-{compactar(toks)}'
    else:
        base = f'{pref}-{slug(idr)}' + ('' if idr.upper() not in ('S/N', 'PASILLO') else '-' + slug(inq)[:18])
    idl = base
    if idl in ids_usados: idl = f'{base}-B{int(num(blk))}'
    n = 2
    while idl in ids_usados: idl = f'{base}-B{int(num(blk))}-{n}'; n += 1
    ids_usados[idl] = True
    notas = txt(r[COL['notas']])
    flag = txt(r[COL['flag']])
    judicial = 'JUDICIAL' in sin_acentos(notas + ' ' + txt(r[COL['precio_viejo']])).upper()
    se_va = 'SE VA' in sin_acentos(notas + ' ' + flag).upper() or 'SE FUE' in sin_acentos(flag).upper()
    vacio = inq_n.startswith('VACIO')
    rubro = txt(r[COL['rubro']]); rubro = '' if rubro in ('-', '') else rubro
    unidades.append({
        'id_local': idl, 'categoria': catn, 'pref': pref, 'block': int(num(blk)), 'sector': f'Block {int(num(blk))}', 'id_planilla': idr, 'tokens': toks,
        'inquilino_planilla': '' if vacio else inq, 'rubro': rubro, 'm2': num(r[COL['m2']]), 'entrepiso': num(r[COL['entrepiso']]), 'vidriera': num(r[COL['vidriera']]),
        'precio_final': num(r[COL['final']]), 'respuesta': num(r[COL['respuesta']]), 'sigue': bool(r[COL['sigue']]) if r[COL['sigue']] not in (None, '') else False,
        'notas': notas, 'flag': flag, 'judicial': judicial, 'se_va': se_va, 'vacio': vacio, 'interno': interno,
        'planta': 'PA' if (pref == 'OFI' and (re.match(r'(?i)^oficina\s*\d', idr) or idr.upper().endswith('PA') or (int(num(blk)) == 3 and toks == {'1', '2', '3'})) and not interno) else ('EXT' if exterior else 'PB'),
        'expensas': expensas.get((compactar(toks), int(num(blk))), 0), 'zonas': [], 'cobranzas': defaultdict(list), 'origen': 'planilla alquileres',
    })
print(f'Unidades en "Locales Comerciales": {len(unidades)}')

# ============================================================ 2. Cobranzas
MESES = {'MAYO': '2026-05', 'JUNIO': '2026-06', 'JULIO': '2026-07', 'AGOSTO': '2026-08'}
ALIAS_NOMBRE = {'CENCOSUD': 'LOC-S1', 'VILLA GLORIA': 'BOL-32-33', 'ACONQUIJA': 'BOL-65-66'}
ALIAS_ID = {'PN2': 'PN1', 'CARTEL KING PREDIO NORTE': 'CARTEL'}
wb_cob = load_workbook(XLS_COB, data_only=True)
cobranzas = []   # filas normalizadas
for hoja, periodo in MESES.items():
    if hoja not in wb_cob.sheetnames: print(f'ATENCIÓN: falta la hoja {hoja} en cobranzas'); continue
    h = wb_cob[hoja]
    filas = [list(r) for r in h.iter_rows(values_only=True)]
    # Celdas combinadas (un "5" que abarca tres filas): openpyxl devuelve None fuera de la celda ancla
    for rng in h.merged_cells.ranges:
        v = h.cell(rng.min_row, rng.min_col).value
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                if rr - 1 < len(filas) and cc - 1 < len(filas[rr - 1]) and filas[rr - 1][cc - 1] is None: filas[rr - 1][cc - 1] = v
    ih = next((i for i, f in enumerate(filas) if any(txt(c).upper().startswith('NUMERO LOCAL') for c in f)), None)
    if ih is None: print(f'ATENCIÓN: no encuentro el encabezado NUMERO LOCAL en {hoja}'); continue
    encc = [txt(c).upper() for c in filas[ih]]
    def col(pref):
        for j, c in enumerate(encc):
            if c.startswith(pref): return j
        return None
    c_id, c_nom, c_tot, c_tr, c_ef, c_ch, c_ret, c_rec = col('NUMERO'), col('SUBCON'), col('TOTAL'), col('TRANS'), col('EFECTIVO'), col('CHE'), col('RETENC'), col('RECIBIDO')
    faltan = [n for n, c in (('TOTAL', c_tot), ('TRANS', c_tr), ('EFECTIVO', c_ef), ('CHE', c_ch), ('RECIBIDO', c_rec)) if c is None]
    if faltan: print(f'ATENCIÓN: en {hoja} faltan columnas {faltan}. No importo ese mes.'); continue
    for f in filas[ih + 1:]:
        idc, nom = limpiar_id(f[c_id]), txt(f[c_nom])
        if not nom or nom.upper() in ('TRANSF', 'EFECT', 'CHEQ', 'RET', 'SALDO A COBRAR'): continue
        total = num(f[c_tot])
        pagos = []
        for c, medio in ((c_tr, 'TRANSFERENCIA'), (c_ef, 'EFECTIVO'), (c_ch, 'CHEQUE')):
            m = num(f[c])
            if m: pagos.append({'medio': medio, 'monto': m, 'fecha': fecha(f[c + 1]) if c + 1 < len(f) else ''})
        ret = num(f[c_ret]) if c_ret is not None else 0
        if ret: pagos.append({'medio': 'RETENCIÓN', 'monto': ret, 'fecha': next((p['fecha'] for p in pagos if p['fecha']), '')})
        cobranzas.append({'hoja': hoja, 'periodo': periodo, 'id': idc, 'nombre': nom, 'total': total, 'pagos': pagos, 'recibido': num(f[c_rec])})
print(f'Filas de cobranzas: {len(cobranzas)}')

def candidatos_por_id(idc):
    idn = sin_acentos(idc).upper().strip()
    for k, v in ALIAS_ID.items(): idn = idn.replace(k, v)
    if not idn: return [], set()
    tk = tokens(idn)
    es_of = 'OFICINA' in idn or 'PA' in idn.split()
    out = []
    for u in unidades:
        if not tk:
            if slug(idn) in u['id_local'] or slug(idn) in slug(u['id_planilla']): out.append(u)
            continue
        if not (u['tokens'] & tk): continue
        if es_of and u['pref'] != 'OFI': continue
        if not es_of and u['pref'] == 'OFI': continue
        out.append(u)
    return out, tk

revisar = []
for c in cobranzas:
    if c['total'] < 0: revisar.append(('Cobranza con total negativo (ajuste), no se cargó', f"{c['hoja']} · {c['id']} · {c['nombre']} · {c['total']}")); continue
    dest = None; conf = ''
    for k, v in ALIAS_NOMBRE.items():
        if k in sin_acentos(c['nombre']).upper(): dest = next((u for u in unidades if u['id_local'] == v), None); conf = 'alias por nombre'
    if not dest:
        cands, tk = candidatos_por_id(c['id'])
        if cands:
            def puntaje(u):
                p = 0
                if u['tokens'] == tk: p += 3
                elif tk <= u['tokens'] or u['tokens'] <= tk: p += 2
                else: p += 1
                s = similitud(c['nombre'], u['inquilino_planilla'])
                if s >= .5: p += 3
                elif s >= .3: p += 1
                return p, s
            cands.sort(key=lambda u: puntaje(u), reverse=True)
            mejor = cands[0]; p, s = puntaje(mejor)
            if len(cands) == 1 or p >= 4 or (p >= 3 and puntaje(cands[1])[0] < p):
                dest = mejor; conf = f'número{" + nombre" if s >= .5 else ""}'
                if s < .3 and mejor['inquilino_planilla']: revisar.append(('Cobranza asignada por número pero el nombre no coincide con la planilla de alquileres (¿cambió el inquilino?)', f"{c['hoja']} · {c['id']} · {c['nombre']} → {mejor['id_local']} ({mejor['inquilino_planilla']})"))
            else:
                revisar.append(('Cobranza ambigua: varias unidades con ese número y ningún nombre coincide. Se creó una unidad aparte', f"{c['hoja']} · {c['id']} · {c['nombre']} · candidatas: " + ', '.join(u['id_local'] for u in cands[:4])))
    if not dest:
        # Unidad nueva desde cobranzas (no existe en la planilla de alquileres)
        tk = tokens(c['id']); idn = 'COB-' + (compactar(tk) if tk else slug(c['id'] or c['nombre'])[:20])
        dest = next((u for u in unidades if u['id_local'] == idn), None)
        if not dest:
            dest = {'id_local': idn, 'categoria': 'Otro', 'pref': 'COB', 'block': 0, 'sector': '', 'id_planilla': c['id'], 'tokens': tk, 'inquilino_planilla': c['nombre'],
                    'rubro': '', 'm2': 0, 'entrepiso': 0, 'vidriera': 0, 'precio_final': 0, 'respuesta': 0, 'sigue': True, 'notas': 'Creada desde cobranzas: no figura en la planilla de alquileres',
                    'flag': '', 'judicial': False, 'se_va': False, 'vacio': False, 'interno': False, 'planta': 'PB', 'expensas': 0, 'zonas': [], 'cobranzas': defaultdict(list), 'origen': 'cobranzas'}
            unidades.append(dest)
            revisar.append(('Unidad que cobra pero no está en la planilla de alquileres (se creó como "Otro")', f"{c['id']} · {c['nombre']}"))
        conf = 'nueva desde cobranzas'
    c['unidad'] = dest['id_local']; c['confianza'] = conf
    dest['cobranzas'][c['periodo']].append(c)

# ============================================================ 3. Zonas del plano
zonas = [{'id': z['id_local'], 'tokens': tokens(z['etiqueta_plano']), 'nombre': z['nombre_en_plano'], 'sector': z['sector_nombre'], 'area': z['caja'][2] * z['caja'][3],
          'norte': z['caja'][1] < 322, 'unidad': None, 'conf': 0} for z in PLANO['locales']]   # 'norte': hilera de boleterías frente a los andenes
pares = []
for z in zonas:
    for u in unidades:
        if u['pref'] == 'GON' or u['planta'] != 'PB' or not u['tokens']: continue
        zt, ut = z['tokens'], u['tokens']
        base = lambda ts: {re.sub(r'[A-D]$', '', t) for t in ts}
        directo = bool(zt & ut)
        # Coincidencia por número base solo cuando la UNIDAD tiene letras y la zona no (61 ↔ 61A-61B).
        # Al revés (zona 10a&b vs unidad 10) no vale: la letra en el plano indica otra subdivisión.
        letras = lambda ts: any(re.search(r'\d[A-D]$', t) for t in ts)
        if not directo and not (letras(ut) and not letras(zt) and (base(zt) & base(ut))): continue
        if not directo: zt, ut = base(zt), base(ut)
        frac = len(zt & ut) / len(zt | ut)
        if zt == ut: p = 5
        elif zt <= ut: p = 2 + frac          # la zona es una parte de la unidad (41A dentro de 41A-41B)
        elif ut <= zt: p = 1 + frac          # la unidad es una parte de la zona (zona 11-15 vs unidad 12): débil
        else: p = 1 + frac
        if not directo: p -= 0.5
        if z['sector'] == u['sector']: p += 1
        # Hilera norte (frente a andenes) = boleterías; hilera sur = locales
        if z['sector'] in ('Block 1', 'Block 2', 'Block 3', 'Block 4'):   # solo ahí conviven boleterías y locales
            if u['pref'] == 'BOL': p += 1.5 if z['norte'] else -1
            elif u['pref'] in ('LOC', 'ENC', 'GON'): p += 1 if not z['norte'] else -1
        if u['pref'] in ('DEP', 'OFI') and (z['nombre'] == '' or 'DEP' in z['nombre'].upper() or u['interno']): p += 1
        nombres = [u['inquilino_planilla'], u['rubro']] + [c['nombre'] for per in u['cobranzas'].values() for c in per]
        if z['nombre'] and z['nombre'] != 'VACÍO' and max((similitud(z['nombre'], n) for n in nombres if n), default=0) >= .5: p += 2
        if z['nombre'] == 'VACÍO' and u['vacio']: p += 1
        pares.append((p, z, u))
pares.sort(key=lambda x: -x[0])
for p, z, u in pares:
    if z['unidad'] or p < 2.5: continue
    z['unidad'] = u['id_local']; z['conf'] = p; u['zonas'].append(z['id'])
for u in unidades:
    if u['zonas']:
        secs = {z['sector'] for z in zonas if z['unidad'] == u['id_local']}
        if u['origen'] == 'cobranzas' and secs: u['sector'] = sorted(secs)[0]
sin_zona = [u for u in unidades if not u['zonas'] and u['pref'] not in ('GON', 'OFI') and u['planta'] == 'PB']
zonas_libres = [z for z in zonas if not z['unidad']]
print(f'Zonas del plano: {len(zonas)} · asignadas: {len(zonas) - len(zonas_libres)} · sin unidad: {len(zonas_libres)}')
print(f'Unidades PB sin zona en el plano: {len(sin_zona)}')

# ============================================================ 4. Estado, contrato, cuotas
ultimo_mes = max(MESES.values())
for u in unidades:
    cob = u['cobranzas']
    ult = None
    for per in sorted(cob.keys(), reverse=True):
        if cob[per]: ult = cob[per]; break
    if u['interno']:
        u['estado'] = 'LIBRE'; u['activo'] = False; u['contrato'] = None
        u['obs'] = f"Uso interno / organismo: {u['inquilino_planilla']} (sin alquiler)"
    elif u['judicial']:
        u['estado'] = 'JUDICIAL'; u['activo'] = True
        u['contrato'] = {'inquilino': u['inquilino_planilla'], 'monto': u['respuesta'] or u['precio_final'], 'obs': 'GESTIÓN JUDICIAL según planilla de alquileres. ' + u['notas']}
        u['obs'] = 'En gestión judicial'
    elif u['vacio'] or not u['inquilino_planilla']:
        u['estado'] = 'LIBRE'; u['activo'] = True; u['contrato'] = None; u['obs'] = ''
    elif ult or u['sigue'] or (cob and not u['se_va']):
        u['estado'] = 'ALQUILADO'; u['activo'] = True
        inq = ult[-1]['nombre'].title() if ult else u['inquilino_planilla']
        monto = sum(c['total'] for c in ult) if ult else (u['respuesta'] or u['precio_final'])
        obs = []
        if ult: obs.append(f"Monto = total facturado {ult[0]['periodo']} todo concepto c/IVA (planilla de cobranzas)")
        if u['respuesta']: obs.append(f"Precio final acordado s/planilla alquileres: ${u['respuesta']:,.0f}".replace(',', '.'))
        if u['expensas']: obs.append(f"Expensas s/planilla (Hoja 11): ${u['expensas']:,.0f} (verificar si ya están dentro del todo concepto)".replace(',', '.'))
        if u['se_va']: obs.append('SE VA según planilla de alquileres')
        if u['notas']: obs.append('Nota planilla: ' + u['notas'])
        if ult and similitud(ult[-1]['nombre'], u['inquilino_planilla']) < .3 and u['inquilino_planilla']: obs.append(f"En la planilla de alquileres figura como: {u['inquilino_planilla']}")
        u['contrato'] = {'inquilino': inq, 'monto': monto, 'obs': ' · '.join(obs)}
        u['obs'] = 'Se va (s/planilla)' if u['se_va'] else ''
    else:
        u['estado'] = 'LIBRE'; u['activo'] = True; u['contrato'] = None
        u['obs'] = f"Ex inquilino: {u['inquilino_planilla']}" + (' (se fue)' if u['se_va'] else ' (no sigue)')

    cn = {'BOL': 'Boletería', 'LOC': 'Local', 'DEP': 'Depósito', 'GON': 'Góndola', 'OFI': 'Oficina', 'ENC': 'Encomiendas', 'CD': 'Corta distancia', 'COB': 'Local', 'OTR': 'Unidad'}[u['pref']]
    numero = compactar(u['tokens']) if u['tokens'] else u['id_planilla']
    if u['pref'] == 'OFI' and u['id_planilla'].upper().startswith('OFICINA'): numero = u['id_planilla'].split()[-1]
    if u['pref'] == 'OFI' and u['id_planilla'].upper().startswith('LOCAL'): numero = u['id_planilla'].replace('Local', '').strip()
    u['nombre'] = f'{cn} {numero}' if u['tokens'] else f"{cn} · {u['id_planilla']}"
    if u['interno']: u['nombre'] += f" · {u['inquilino_planilla']}"
    extras = []
    if u['entrepiso']: extras.append(f"Entrepiso {u['entrepiso']:g} m²")
    if u['vidriera']: extras.append(f"Vidriera {u['vidriera']:g} m²")
    if u['origen'] == 'planilla alquileres': extras.append(f"Id en planilla alquileres: {u['id_planilla']} (Block {u['block']})")
    if u['obs']: extras.insert(0, u['obs'])
    u['obs'] = ' · '.join(extras)

# ============================================================ 5. Escribir planilla
LISTAS = {
    'categorias': ['Boletería', 'Local comercial', 'Depósito', 'Góndola', 'Oficina', 'Encomiendas', 'Corta distancia', 'Otro'],
    'rubros': sorted({u['rubro'].strip().capitalize() for u in unidades if u['rubro']} | {'Otro'}, key=lambda s: sin_acentos(s)),
    'sectores': [f'Block {i}' for i in range(1, 9)] + ['Predio norte', 'Planta alta'],
    'indices_ajuste': ['IPC', 'ICL', '% FIJO', 'OTRO'],
    'tipos_falla': ['ELÉCTRICA', 'SANITARIA', 'ESTRUCTURAL', 'PINTURA', 'CERRAJERÍA', 'CLIMATIZACIÓN', 'VIDRIOS / CARPINTERÍA', 'OTRO'],
    'intervinientes': ['Personal INGECO', 'Electricista', 'Plomero', 'Albañil', 'Refrigeración', 'Cerrajero', 'Otro'],
    'medios_pago': ['TRANSFERENCIA', 'EFECTIVO', 'CHEQUE', 'RETENCIÓN', 'DÉBITO AUTOMÁTICO'],
}
for u in unidades:
    if u['rubro']: u['rubro'] = u['rubro'].strip().capitalize()
    if u['planta'] == 'EXT': u['sector'] = 'Predio norte'
    if u['planta'] == 'PA' and not u['sector'].startswith('Block'): u['sector'] = 'Planta alta'

F_TIT = Font(bold=True, color='FFFFFF', name='Arial'); R_TIT = PatternFill('solid', fgColor='0F766E')
R_FUENTE = PatternFill('solid', fgColor='E2E8F0'); R_OJO = PatternFill('solid', fgColor='FEF3C7')
BORDE = Border(*(Side(style='thin', color='CBD5E1'),) * 4); F_NORMAL = Font(name='Arial', size=10)
def encabezar(ws, cols, anchos):
    for j, c in enumerate(cols, 1):
        cel = ws.cell(row=1, column=j, value=c); cel.font = F_TIT; cel.fill = R_TIT; cel.border = BORDE
        cel.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = anchos[j - 1]
    ws.row_dimensions[1].height = 30; ws.freeze_panes = 'A2'
def fila(ws, i, valores, fuente_hasta=0, ojo=False):
    for j, v in enumerate(valores, 1):
        c = ws.cell(row=i, column=j, value=v); c.border = BORDE; c.font = F_NORMAL
        if j <= fuente_hasta: c.fill = R_FUENTE
        elif ojo: c.fill = R_OJO
def validar(ws, col, lista, desde, hasta):
    k = list(LISTAS.keys()).index(lista) + 1; L = get_column_letter(k)
    dv = DataValidation(type='list', formula1=f"='LISTAS'!${L}$2:${L}${len(LISTAS[lista]) + 40}", allow_blank=True, showErrorMessage=True, errorTitle='Valor no permitido', error='Elegí una opción de la lista (podés agregar opciones en LISTAS).')
    ws.add_data_validation(dv); dv.add(f'{col}{desde}:{col}{hasta}')
def validar_fija(ws, col, ops, desde, hasta):
    dv = DataValidation(type='list', formula1='"' + ','.join(ops) + '"', allow_blank=True, showErrorMessage=True); ws.add_data_validation(dv); dv.add(f'{col}{desde}:{col}{hasta}')

wb = Workbook(); ws = wb.active; ws.title = 'INSTRUCCIONES'; ws.column_dimensions['A'].width = 125
texto = [
    ('Relevamiento de locales · Terminal de Ómnibus · INGECO S.A.', True),
    (f'Generada el {HOY:%d/%m/%Y} cruzando: planilla "Alquileres Terminal" (lista de unidades, superficies, precios), planilla de COBRANZAS mayo–agosto 2026 (cuotas y pagos) y el plano CONFORME A OBRA 2026 (zonas).', False),
    ('', False), ('CÓMO ESTÁ ARMADA', True),
    ('• LOCALES: una fila por unidad (boletería, local, depósito, góndola, oficina, encomiendas, predio norte). Las columnas grises salen de las planillas fuente. Completar/corregir: nombre, m², rubro, estado, se_alquila, zonas_plano.', False),
    ('• CONTRATOS: una fila por unidad ocupada. Inquilino = razón social de la última cobranza. Monto = último total facturado (todo concepto c/IVA). FALTAN fechas de inicio/fin, CUIT, depósito, índice: completar.', False),
    ('• CUOTAS y PAGOS: mayo a agosto 2026, tal como están en la planilla de cobranzas. Concepto ALQUILER = "todo concepto". Se pueden corregir montos o fechas acá.', False),
    ('• CRUCE: cómo quedó emparejada cada zona del plano con cada unidad, y cada fila de cobranzas con cada unidad. Solo informativo.', False),
    ('• REVISAR: la lista de dudas que quedaron del cruce automático (nombres que no coinciden, unidades que cobran pero no figuran, zonas del plano sin unidad). Es lo primero que conviene mirar.', False),
    ('• LISTAS: opciones de los desplegables; se pueden agregar valores al final de cada columna.', False),
    ('', False), ('REGLAS', True),
    ('• Los ids (BOL-1-2, LOC-501, GON-3, OFI-7...) identifican la unidad en todo el sistema: no cambiarlos. El número solo se repite entre categorías.', False),
    ('• zonas_plano: ids de las zonas del plano (L-XX) que ocupa la unidad, separados por ";". Vacío = no se dibuja en el mapa (góndolas del hall, oficinas de planta alta, predio norte).', False),
    ('• estado: LIBRE / ALQUILADO / REFACCION / JUDICIAL. Una unidad con fila en CONTRATOS tiene que estar ALQUILADO o JUDICIAL.', False),
    ('• se_alquila = NO: uso interno u organismo público; queda en el mapa pero no cuenta para ocupación ni deuda.', False),
    ('• Fechas dd/mm/aaaa. Montos sin puntos ni "$".', False),
    ('', False), ('QUÉ PASA DESPUÉS', True),
    ('Con la planilla corregida se corre herramientas/generar-datos.py y la app pasa a mostrar estos datos. Desde ahí, los meses siguientes se cargan en la app ("+ Cuotas del mes" y "Registrar pago").', False),
]
for i, (t, b) in enumerate(texto, 1):
    c = ws.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=True, vertical='top'); c.font = Font(name='Arial', bold=b, size=12 if i == 1 else 10)

# LOCALES
ws = wb.create_sheet('LOCALES')
colsL = ['id_local', 'categoria', 'planta', 'sector', 'id_planilla_alquileres', 'inquilino_planilla', 'nombre_en_plano', 'nombre', 'zonas_plano', 'm2', 'rubro', 'estado', 'se_alquila', 'observaciones']
encabezar(ws, colsL, [14, 15, 7, 12, 22, 30, 24, 22, 20, 8, 24, 12, 10, 60])
orden = {'Block 1': 1, 'Block 2': 2, 'Block 3': 3, 'Block 4': 4, 'Block 5': 5, 'Block 6': 6, 'Block 7': 7, 'Block 8': 8}
unidades.sort(key=lambda u: (orden.get(u['sector'], 9), ['BOL', 'LOC', 'ENC', 'DEP', 'GON', 'OFI', 'CD', 'COB', 'OTR'].index(u['pref']), compactar(u['tokens']).zfill(6) if u['tokens'] else 'zz'))
i = 2
for u in unidades:
    nombre_plano = ' / '.join(z['nombre'] for z in zonas if z['unidad'] == u['id_local'] and z['nombre'])
    ojo = u['origen'] == 'cobranzas' or (u['planta'] == 'PB' and not u['zonas'] and u['pref'] not in ('GON', 'OFI'))
    fila(ws, i, [u['id_local'], u['categoria'], u['planta'], u['sector'], u['id_planilla'], u['inquilino_planilla'], nombre_plano, u['nombre'], ';'.join(u['zonas']),
                 u['m2'] or None, u['rubro'], u['estado'], 'SI' if u['activo'] else 'NO', u['obs']], fuente_hasta=7, ojo=ojo)
    i += 1
nL = i - 1
validar(ws, 'B', 'categorias', 2, nL + 300); validar(ws, 'D', 'sectores', 2, nL + 300); validar(ws, 'K', 'rubros', 2, nL + 300)
validar_fija(ws, 'C', ['PB', 'PA', 'EXT'], 2, nL + 300); validar_fija(ws, 'L', ['LIBRE', 'ALQUILADO', 'REFACCION', 'JUDICIAL'], 2, nL + 300); validar_fija(ws, 'M', ['SI', 'NO'], 2, nL + 300)
ws.auto_filter.ref = f'A1:{get_column_letter(len(colsL))}{nL}'

# CONTRATOS
ws = wb.create_sheet('CONTRATOS')
colsC = ['id_local', 'inquilino', 'cuit', 'contacto', 'fecha_inicio', 'fecha_fin', 'monto_alquiler', 'indice_ajuste', 'periodicidad_meses', 'proxima_fecha_ajuste', 'deposito', 'expensas_mensuales', 'observaciones']
encabezar(ws, colsC, [14, 36, 15, 26, 13, 13, 15, 12, 12, 14, 13, 14, 90])
i = 2
for u in unidades:
    if not u['contrato']: continue
    c = u['contrato']
    fila(ws, i, [u['id_local'], c['inquilino'], None, None, None, None, round(c['monto']) or None, 'IPC', 3, None, None, 0, c['obs']], fuente_hasta=2)
    for j in (5, 6, 10): ws.cell(row=i, column=j).number_format = 'DD/MM/YYYY'
    for j in (7, 11, 12): ws.cell(row=i, column=j).number_format = '#,##0'
    i += 1
nC = i - 1
validar(ws, 'H', 'indices_ajuste', 2, nC + 300); validar_fija(ws, 'I', ['1', '2', '3', '4', '6', '12'], 2, nC + 300)
ws.auto_filter.ref = f'A1:{get_column_letter(len(colsC))}{nC}'

# CUOTAS y PAGOS
wq = wb.create_sheet('CUOTAS'); colsQ = ['id_local', 'periodo', 'concepto', 'monto', 'vencimiento', 'observaciones']
encabezar(wq, colsQ, [14, 10, 12, 15, 13, 70])
wp = wb.create_sheet('PAGOS'); colsP = ['id_local', 'periodo', 'concepto', 'fecha', 'monto', 'medio', 'comprobante', 'observaciones']
encabezar(wp, colsP, [14, 10, 12, 13, 15, 16, 14, 60])
iq = ip = 2; nq = npg = 0
for u in unidades:
    for per in sorted(u['cobranzas'].keys()):
        filas_c = u['cobranzas'][per]
        total = sum(c['total'] for c in filas_c)
        if total <= 0 and not any(c['pagos'] for c in filas_c): continue
        obs = 'Todo concepto c/IVA · ' + ' + '.join(f"{c['nombre']} ({c['id']})" for c in filas_c)
        y, m = per.split('-')
        fila(wq, iq, [u['id_local'], per, 'ALQUILER', round(total, 2), datetime.date(int(y), int(m), 10), obs], fuente_hasta=3)
        wq.cell(row=iq, column=5).number_format = 'DD/MM/YYYY'; wq.cell(row=iq, column=4).number_format = '#,##0.00'; iq += 1; nq += 1
        for c in filas_c:
            for p in c['pagos']:
                f = datetime.date.fromisoformat(p['fecha']) if p['fecha'] else None
                fila(wp, ip, [u['id_local'], per, 'ALQUILER', f, round(p['monto'], 2), p['medio'], '', c['nombre'] + ('' if f else ' · SIN FECHA en la planilla')], fuente_hasta=3, ojo=not f)
                wp.cell(row=ip, column=4).number_format = 'DD/MM/YYYY'; wp.cell(row=ip, column=5).number_format = '#,##0.00'; ip += 1; npg += 1
            suma = sum(p['monto'] for p in c['pagos'])
            if abs(suma - c['recibido']) > 1: revisar.append(('La suma de pagos no coincide con la columna RECIBIDO', f"{c['hoja']} · {c['id']} · {c['nombre']}: pagos {suma:,.0f} vs recibido {c['recibido']:,.0f}"))
validar_fija(wq, 'C', ['ALQUILER', 'EXPENSAS'], 2, iq + 500); validar_fija(wp, 'C', ['ALQUILER', 'EXPENSAS'], 2, ip + 800); validar(wp, 'F', 'medios_pago', 2, ip + 800)
wq.auto_filter.ref = f'A1:F{iq - 1}'; wp.auto_filter.ref = f'A1:H{ip - 1}'

# MANTENIMIENTO (vacía, misma estructura de antes)
ws = wb.create_sheet('MANTENIMIENTO'); colsM = ['id_local', 'fecha_reporte', 'tipo_falla', 'descripcion', 'prioridad', 'estado', 'quien_intervino', 'fecha_resolucion', 'costo', 'observaciones']
encabezar(ws, colsM, [14, 13, 20, 40, 11, 12, 20, 14, 12, 40])
validar(ws, 'C', 'tipos_falla', 2, 500); validar_fija(ws, 'E', ['ALTA', 'MEDIA', 'BAJA'], 2, 500); validar_fija(ws, 'F', ['PENDIENTE', 'EN CURSO', 'RESUELTO'], 2, 500); validar(ws, 'G', 'intervinientes', 2, 500)

# LISTAS
ws = wb.create_sheet('LISTAS'); encabezar(ws, list(LISTAS.keys()), [18, 30, 14, 14, 24, 20, 20])
for j, k in enumerate(LISTAS.keys(), 1):
    for i2, v in enumerate(LISTAS[k], 2): ws.cell(row=i2, column=j, value=v).border = BORDE

# CRUCE
ws = wb.create_sheet('CRUCE'); encabezar(ws, ['zona_plano', 'etiqueta_plano', 'nombre_en_plano', 'sector_plano', 'unidad_asignada', 'inquilino_unidad', 'puntaje'], [12, 14, 26, 10, 16, 34, 8])
i = 2
for z in sorted(zonas, key=lambda z: (z['unidad'] is not None, z['id'])):
    u = next((x for x in unidades if x['id_local'] == z['unidad']), None)
    et = next(l['etiqueta_plano'] for l in PLANO['locales'] if l['id_local'] == z['id'])
    fila(ws, i, [z['id'], et, z['nombre'], z['sector'], z['unidad'] or '', (u['contrato']['inquilino'] if u and u['contrato'] else (u['inquilino_planilla'] if u else '')), z['conf'] or None], ojo=not z['unidad']); i += 1
i += 1; fila(ws, i, ['COBRANZAS → UNIDAD'], 0); ws.cell(row=i, column=1).font = Font(name='Arial', bold=True); i += 1
fila(ws, i, ['hoja', 'nro local (cobranzas)', 'subconcesionario', 'total', 'unidad', 'confianza'], 0); i += 1
for c in cobranzas:
    fila(ws, i, [c['hoja'], c['id'], c['nombre'], c['total'], c.get('unidad', ''), c.get('confianza', 'no cargada')]); i += 1

# REVISAR
ws = wb.create_sheet('REVISAR'); encabezar(ws, ['qué revisar', 'detalle'], [70, 110])
for z in zonas_libres: revisar.append(('Zona del plano sin unidad asignada (no se pinta en el mapa)', f"{z['id']} · rótulo \"{z['nombre']}\" · {z['sector']}"))
for u in sin_zona: revisar.append(('Unidad de planta baja sin zona en el plano (no se ve en el mapa)', f"{u['id_local']} · {u['nombre']} · {u['inquilino_planilla']} · {u['sector']}"))
vistos = set(); i = 2
for q, d in revisar:
    if (q, d) in vistos: continue
    vistos.add((q, d)); fila(ws, i, [q, d]); i += 1
ws.auto_filter.ref = f'A1:B{i - 1}'
wb.move_sheet('REVISAR', offset=-(len(wb.sheetnames) - 2))

os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
wb.save(SALIDA)
print(f'OK → {SALIDA}')
print(f'  LOCALES {nL - 1} · CONTRATOS {nC - 1} · CUOTAS {nq} · PAGOS {npg} · REVISAR {i - 2}')
est = defaultdict(int)
for u in unidades: est[u['estado'] if u['activo'] else 'USO INTERNO'] += 1
print('  estados:', dict(est))
